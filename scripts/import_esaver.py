"""One-time, read-only import of the operator's schedule. Never saves the workbook."""
import argparse
import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import openpyxl
from src.state import write_json
DATA_DIR = Path(__file__).resolve().parents[1] / "data"
PROMOTIONS_FILE = DATA_DIR / "esaver_promotions.json"
REGISTRATIONS_FILE = DATA_DIR / "esaver_registrations.json"


def import_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    sheet = workbook["Sheet1"]
    expected = ["Start", "End", "No. of Days", "Previous Valid Promo End", "Gap Days", "Benchmark Date", "Not for Month", "0.2m HKD", "1m HKD", "5m HKD", "25k USD", "Reg Deadline", "Reg"]
    if [c.value for c in sheet[1]] != expected:
        raise ValueError("Workbook columns changed; refusing ambiguous import")
    promotions, registrations, members = [], {}, []
    def iso(value):
        return value.date().isoformat() if isinstance(value, datetime) else None
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, 14)]
        if not values[0]:
            continue
        start, end = iso(values[0]), iso(values[1])
        key = start[:7]
        excluded = []
        for token in str(values[6] or "").split(","):
            if token.strip():
                month = int(token)
                year = int(key[:4]) - (1 if month > int(key[5:]) else 0)
                excluded.append(f"{year}-{month:02d}")
        rate_keys = ["hkd_200k", "hkd_1m", "hkd_5m", "usd_25k"]
        rates = {k: round(values[i + 7] * 100, 5) if isinstance(values[i + 7], (float, int)) else None for i, k in enumerate(rate_keys)}
        offer = {"id": key, "promo_month": key, "reward_start": start, "reward_end": end,
                 "days": (datetime.fromisoformat(end) - datetime.fromisoformat(start)).days + 1,
                 "previous_valid_end": iso(values[3]), "gap_days": values[4], "benchmark_date": iso(values[5]),
                 "excluded_months": excluded, "rates": rates, "reg_end": iso(values[11]),
                 "credit_date": None, "source_kind": "excel", "source_label": "esaver schedule.xlsx",
                 "excel_row": row, "excel_formulas": {formula_book["Sheet1"].cell(row, col).coordinate: formula_book["Sheet1"].cell(row, col).value for col in [3,4,5] if formula_book["Sheet1"].cell(row, col).data_type == "f"}}
        if values[2] != offer["days"]:
            raise ValueError(f"Row {row}: cached duration does not match raw dates")
        promotions.append(offer)
        if values[12]:
            member = str(values[12]).strip()
            members.append(member)
            registrations[f"{key}|{member}"] = {"promo_id": key, "member": member, "status": "registered", "registered_on": None, "source": "excel", "excel_row": row}
    if PROMOTIONS_FILE.exists() or REGISTRATIONS_FILE.exists():
        raise FileExistsError("Import already exists; refusing to overwrite family records")
    write_json(PROMOTIONS_FILE, {"promotions": sorted(promotions,key=lambda p:p["reward_start"]),
               "import": {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "rows": len(promotions)}})
    write_json(REGISTRATIONS_FILE, {"members": list(dict.fromkeys(members)), "registrations": registrations, "version": 1, "sync_errors": []})
    print(json.dumps({"promotions": len(promotions), "registrations": len(registrations), "members": list(dict.fromkeys(members))}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    import_workbook(parser.parse_args().path)
