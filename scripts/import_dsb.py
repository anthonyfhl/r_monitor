"""Read the supplied balance workbook without executing its contents or saving it.

Only numeric literals, cell references, SUM and arithmetic are accepted. Original
cached results are preserved for comparison, never used as the calculated totals.
"""
import argparse
import ast
import calendar
from datetime import datetime, date
import hashlib
import json
import operator
from pathlib import Path
import re
import sys

import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src import dsb
from src.state import read_json, write_json


def numeric(sheet, cell, trail=()):
    if cell in trail:
        raise ValueError('Workbook arithmetic has a circular reference: '+cell)
    value=sheet[cell].value
    if type(value) in (int,float): return value
    if not isinstance(value,str) or not value.startswith('='):
        raise ValueError('Missing numeric input: '+cell)
    formula=value[1:].replace('$','')
    def total(match):
        return str(sum(numeric(sheet,c.coordinate,trail+(cell,)) for row in sheet[match.group(1)] for c in row))
    formula=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',total,formula)
    formula=re.sub(r'([0-9.]+)%',r'(\1/100)',formula)
    formula=re.sub(r'\b[A-Z]+[0-9]+\b',lambda m:str(numeric(sheet,m.group(),trail+(cell,))),formula)
    ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
    def compute(node):
        if isinstance(node,ast.Constant) and type(node.value) in (int,float): return node.value
        if isinstance(node,ast.BinOp) and type(node.op) in ops:return ops[type(node.op)](compute(node.left),compute(node.right))
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):return compute(node.operand)*(1 if isinstance(node.op,ast.UAdd) else -1)
        raise ValueError('Unsupported workbook expression at '+cell)
    return compute(ast.parse(formula,mode='eval').body)


def extract(path):
    book=openpyxl.load_workbook(path,data_only=False)
    cached=openpyxl.load_workbook(path,data_only=True)
    sheet, old=book['Sheet1'],cached['Sheet1']
    blocks=[(1,2,3,'B36','B37','B38'),(5,6,7,'F36','F37','F38'),(9,10,11,'H36','H37','H38'),
      (13,14,15,'M36','M37','M38'),(17,18,19,'M36','M37','M38'),(21,22,23,'M36','M37','M38'),
      (25,26,27,'M36','M37','M38'),(29,30,31,'M36','M37','M38'),(33,34,35,'M36','M37','M38'),(37,38,39,'AL36','AL42','AL43')]
    records=dsb.empty_records();rules={};audit=[]
    for dc,bc,ic,rc,vc,basic in blocks:
        dates=[old.cell(r,dc).value for r in range(2,33) if isinstance(old.cell(r,dc).value,datetime)]
        if not dates:raise ValueError('Missing workbook date block')
        month=dates[0].strftime('%Y-%m'); days=calendar.monthrange(dates[0].year,dates[0].month)[1]
        values={}
        for row in range(2,33):
            day=old.cell(row,dc).value; cell=sheet.cell(row,bc)
            if not isinstance(day,datetime) or cell.value is None:continue
            if day.strftime('%Y-%m')!=month:raise ValueError('Workbook date crosses month')
            amount=numeric(sheet,cell.coordinate)
            if amount<0:raise ValueError('Negative balance in source')
            values[day.date().isoformat()]=amount
        core=numeric(sheet,rc)*100;vip=numeric(sheet,vc)*100;base=numeric(sheet,basic)*100
        raw=sum(values.values())*core/100/365
        revision='excel-'+month
        rules[revision]={'id':revision,'revision':revision,'reward_start':month+'-01','reward_end':f'{month}-{days:02}',
          'cap':6000,'day_basis':365,'balance_basis':'entered','legacy_core_rate':core,'legacy_vip_rate':vip,'legacy_basic_rate':base,
          'source_kind':'excel','scope_note':'按原表每日結餘及固定年息重算；歷史銀行條款未重新核對。', 'source_label':path.name}
        tasks={'payroll':False,'debit':0,'credit':0,'fx':0,'fund':False,'stock':0,'vip1':False,'vip2':False}
        records['months'][month]={'offer_id':revision,'rule_revision':revision,'account':'vip','actual':tasks.copy(),'planned':tasks.copy(),
          'basic_rate':base,'costs':dict.fromkeys(dsb.COST_KEYS),'receipts':dict.fromkeys(['core','vip','basic']),
          'planning_date':date(dates[0].year+(dates[0].month==12),dates[0].month%12+1,1).isoformat(),
          'current_balance':None,'eligible_from':month+'-01','movements':[],'terms_confirmed':False,'source':'excel'}
        records['balances'].update(values)
        audit.append({'month':month,'known_days':len(values),'expected_days':days,'recomputed_core_before_cap':raw,'core_after_cap':min(6000,raw),
          'original_cached_sum':old.cell(33,ic).value,'source_rate_cell':rc,'complete':len(values)==days})
    records['import']={'filename':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'audit':audit}
    return records,rules


def main():
    parser=argparse.ArgumentParser();parser.add_argument('path',type=Path);parser.add_argument('--check-only',action='store_true');args=parser.parse_args()
    records,rules=extract(args.path)
    if args.check_only:
        print(json.dumps(records['import'],ensure_ascii=False,indent=2));return
    with dsb.record_lock():
        if dsb.RECORDS_FILE.exists():raise FileExistsError('Dah Sing records already exist; refusing overwrite')
        catalog=read_json(dsb.CATALOG_FILE,{'offers':[],'rules':{}})
        catalog['rules'].update(rules)
        # Dates and October participation are explicit operator-provided facts.
        for key,start,end in [('2026-04','2026-04-01','2026-09-30'),('2026-07','2026-07-01','2026-12-31'),('2026-10','2026-10-01','2027-06-30')]:
            if not any(o['id']==key for o in catalog['offers']):
                catalog['offers'].append({'id':key,'reward_start':start,'reward_end':end,'source_kind':'user','revision':None})
        records['enrollments']=[{'offer_id':'2026-10','registered':True,'registered_on':None,'effective_month':'2026-10',
            'reward_end':'2027-06-30','source':'operator confirmed in chat 2026-10-03'}]
        records['version']=1
        write_json(dsb.CATALOG_FILE,catalog);write_json(dsb.RECORDS_FILE,records)
        if dsb.read_records()!=records:raise RuntimeError('Imported records failed verification')
    print(json.dumps({'months':len(records['months']),'daily_balances':len(records['balances']),'sha256':records['import']['sha256']}))


if __name__=='__main__':main()
